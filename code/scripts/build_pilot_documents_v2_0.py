"""Render a frozen Pilot label batch using a retained Word review template."""
import json,sys,copy,zipfile,hashlib,html,re
from pathlib import Path
from collections import Counter
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table

import argparse
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--project-root',type=Path,required=True)
parser.add_argument('--template',type=Path,required=True)
args=parser.parse_args()
P=args.project_root/'data/processed/v2.0/09_pilot_review'
O=args.project_root/'code/results/v2.0/09_pilot_review'
R=args.template
def rows(p):return [json.loads(x) for x in p.read_text(encoding='utf8').splitlines() if x.strip()]
cases=rows(P/'pilot_case_source_records_v2.0.jsonl')
pairs=rows(P/'pilot_label_pairs_v2.0.jsonl')
sources=rows(P/'pilot_evidence_sources_v2.0.jsonl')
EI={e['evidence_id']:e for e in sources}
LABELS={'SUPPORTED':'支持匹配','CONDITIONAL':'条件性匹配','MISMATCH':'有依据不匹配','INDETERMINATE':'无法判定'}
CTX={'routine':'常规使用情境','research':'研究性方案',None:'组成或身份待核实'}
SCOPE={'metastatic_at_diagnosis':'初诊转移性','metastatic_after_diagnosis':'初诊后转移（派生情境）','baseline_unresectable_advanced_current_context_unconfirmed':'初诊不可切除晚期；当前情境未确认'}
DRUGS={'Gemcitabine HCL':'吉西他滨','Albumin-Bound Paclitaxel':'白蛋白结合型紫杉醇','Paclitaxel':'紫杉醇','Fluorouracil':'氟尿嘧啶','Leucovorin':'亚叶酸','Oxaliplatin':'奥沙利铂','Irinotecan HCL':'伊立替康','Liposomal Irinotecan':'脂质体伊立替康','Capecitabine':'卡培他滨','Cisplatin':'顺铂','Mitomycin':'丝裂霉素','Nivolumab':'纳武利尤单抗','Ipilimumab':'伊匹木单抗'}
DRUGS['Nabpaclitaxel']='白蛋白结合型紫杉醇'
TYPES={'pathology_procedure':'病理取材','regimen_start':'疗程开始','ngs_report':'NGS报告','imaging':'影像评估','oncology_assessment':'肿瘤科评估','tumor_marker':'肿瘤标志物','radiation_start':'放疗开始','cancer_diagnosis':'癌症诊断','mmr_ihc':'MMR免疫组化'}
TYPES.update(tumor_marker_collection='肿瘤标志物',mmr_report='MMR报告')
def val(x):return str(x) if x not in ['',None] else '未记录'
def trans(x):
    extra={'Nabpaclitaxel':'白蛋白结合型紫杉醇','External institution':'外院','Internal institution':'本院','Metastasis site unspecified':'转移部位未详','Primary site':'原发部位','Not stated/Indeterminate':'未说明／无法确定','Yes, the Impression states there is evidence of cancer':'评估提示存在癌症证据','Yes, the Impression/Plan states or implies there is evidence of cancer':'评估提示存在癌症证据','Abdomen, Pelvis':'腹部、盆腔','Chest':'胸部','Stage IV':'IV期','Stage III':'III期','Stage II':'II期'}
    if x in extra:return extra[x]
    return {'Male':'男','Female':'女','Yes':'是','No':'否','Adenocarcinoma, NOS':'腺癌 NOS','Ductal carcinoma, NOS':'导管癌 NOS','Pancreatic Cancer':'胰腺癌','Progressing/Worsening/Enlarging':'进展／恶化／增大','Stable/No change':'稳定／无变化','Responding/Improving':'缓解／改善','Surgical pathology':'组织病理','Biopsy':'活检','Cytology':'细胞学','Fine needle aspiration':'细针穿刺'}.get(x,val(x))
def drugs(ds,dates=False):
    ans=[]
    for d in ds:
        n=d['name_raw'].split('(')[0];n='匿名药物（'+n+'）' if d.get('masked') else DRUGS.get(n,n)
        if dates:n+=f" D{d.get('start_day')}至"+(f"D{d['end_day']}" if d.get('end_day') is not None else '截止时点结束日未确认')
        ans.append(n)
    return '＋'.join(ans) or '未记录'
def loc(src):
    if isinstance(src,str):return src
    name=Path(src.get('file','')).name
    name={'pathology_report_level_dataset.csv':'病理表','cancer_level_dataset_index.csv':'癌症表','regimen_cancer_level_dataset.csv':'疗程表','med_onc_note_level_dataset.csv':'肿瘤科记录表','imaging_level_dataset.csv':'影像表','cancer_panel_test_level_dataset.csv':'NGS报告表','tumor_marker_level_dataset.csv':'肿瘤标志物表'}.get(name,name)
    return name+' 行'+str(src.get('logical_row',src.get('physical_line','未记录')))
def facttext(e):
    f=e['facts'];kind=e['kind']
    if kind=='regimen_start':return drugs(f.get('components',[]))
    if kind=='tumor_marker_collection':return f.get('tm_type','')+' '+f.get('tm_num_result','')+' '+f.get('tm_result_units','')+'；参考范围 '+val(f.get('tm_normal_range_lower'))+' 至 '+val(f.get('tm_normal_range_upper'))
    if kind=='imaging':return trans(f.get('image_overall',''))+'；'+ '；'.join(trans(v) for k,v in f.items() if k!='image_overall' and v not in ['',None,[],{}])
    if kind=='oncology_assessment':return trans(f.get('md_ca_status',''))+'；癌种 '+trans(f.get('md_type_ca_cur',''))
    if kind=='ngs_report':return f.get('cpt_genie_sample_id','')+'；'+f.get('cpt_seq_assay_id','')
    if kind=='pathology_procedure':return '；'.join(f"标本{x['specimen_number']} {x['site']} {x['invasive_histology'] or ('浸润性肿瘤 '+trans(x['invasive']))}" for x in f['specimens'])
    return '；'.join(k+'='+str(v) for k,v in f.items() if v not in ['',None,[],{}])

REF=Document(R)
SUMMARY=next(t._tbl for t in REF.tables if len(t.columns)==4)
REVIEW=next(t._tbl for t in REF.tables if len(t.columns)==3)
def font(run,size=10,bold=False,color='000000'):
    run.font.name='宋体';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体' if bold else '宋体')
def para(doc,text='',size=10,bold=False,heading=False,page=False,center=False):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.08
    p.paragraph_format.keep_with_next=bool(heading);p.paragraph_format.widow_control=True
    p.paragraph_format.page_break_before=page
    if heading:p.style='Heading 2';p.paragraph_format.space_before=Pt(9)
    if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(text),size,bold,'365F91' if heading else '000000')
    return p
def celltext(cell,text,bold=False):
    cell.text=''
    for n,line in enumerate(str(text).split('\n')):
        p=cell.paragraphs[0] if n==0 else cell.add_paragraph()
        p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.08
        p.paragraph_format.keep_with_next=False;p.paragraph_format.widow_control=True
        font(p.add_run(line),10,bold)
def shade(cell,color):
    pr=cell._tc.get_or_add_tcPr()
    for old in pr.findall(qn('w:shd')):pr.remove(old)
    x=OxmlElement('w:shd');x.set(qn('w:fill'),color);pr.append(x)
def table(doc,headers,data,widths,summary=False):
    xml=copy.deepcopy(SUMMARY if len(widths)==4 else REVIEW)
    for tr in list(xml.findall(qn('w:tr'))):xml.remove(tr)
    grid=xml.find(qn('w:tblGrid'))
    for x in list(grid):grid.remove(x)
    for width in widths:
        x=OxmlElement('w:gridCol');x.set(qn('w:w'),str(int(width*1440)));grid.append(x)
    doc._body._body.insert(len(doc._body._body)-1,xml);t=Table(xml,doc._body);t.autofit=False
    pr=xml.find(qn('w:tblPr'))
    for tag in ['tblInd','tblBorders','tblCellMar','tblW']:
        for x in list(pr.findall(qn('w:'+tag))):pr.remove(x)
    w=OxmlElement('w:tblW');w.set(qn('w:w'),str(int(sum(widths)*1440)));w.set(qn('w:type'),'dxa');pr.append(w)
    borders=OxmlElement('w:tblBorders')
    for tag in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+tag);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    pr.append(borders);m=OxmlElement('w:tblCellMar')
    for tag,num in [('top',60),('bottom',60),('left',75),('right',75)]:
        x=OxmlElement('w:'+tag);x.set(qn('w:w'),str(num));x.set(qn('w:type'),'dxa');m.append(x)
    pr.append(m)
    allrows=([headers] if headers else [])+data
    for i,values in enumerate(allrows):
        row=t.add_row();rpr=row._tr.get_or_add_trPr();rpr.append(OxmlElement('w:cantSplit'))
        if headers and i==0:rpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,txt) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[j]);celltext(cell,txt,bool(headers and i==0) or (summary and j%2==0))
            shade(cell,'D9E2F3' if headers and i==0 else 'EFEFEF' if summary and j%2==0 else 'FFFFFF')
            if headers and i==0:
                for par in cell.paragraphs:par.paragraph_format.keep_with_next=True
    return t
def new(title,subtitle):
    d=Document(R)
    for e in list(d._body._body):
        if e.tag!=qn('w:sectPr'):d._body._body.remove(e)
    p=para(d,title,18,True,center=True);p.style='Title';p.paragraph_format.space_before=Pt(35)
    borders=OxmlElement('w:pBdr')
    for tag in ['top','left','bottom','right','between']:
        x=OxmlElement('w:'+tag);x.set(qn('w:val'),'nil');borders.append(x)
    p._p.get_or_add_pPr().append(borders)
    para(d,subtitle,14,True,center=True)
    para(d,'AACR GENIE BPC PANC    v2.0',11,center=True)
    para(d,'25 例病例    25 个候选时点    175 条方案标签',11,center=True)
    para(d,'知识截止日 2026年9月12日    编制日期 2026年9月15日',10,center=True)
    return d
def save(d,name):
    # Preserve the original package except explicitly editable body and metadata.
    tmp=O/('_'+name);d.save(tmp)
    with zipfile.ZipFile(R) as z,zipfile.ZipFile(tmp) as generated,zipfile.ZipFile(O/name,'w',zipfile.ZIP_DEFLATED) as out:
        for n in z.namelist():
            b=generated.read(n) if n=='word/document.xml' else z.read(n)
            if n=='docProps/core.xml':
                r=etree.fromstring(b)
                for e in list(r):
                    if etree.QName(e).localname in ['creator','lastModifiedBy','title','subject','description','keywords']:e.text=''
                b=etree.tostring(r,xml_declaration=True,encoding='UTF-8',standalone=True)
            elif n=='docProps/app.xml':
                r=etree.fromstring(b)
                for e in r.iter():
                    if etree.QName(e).localname in ['Company','Manager']:e.text=''
                    if etree.QName(e).localname in ['Pages','Words','Characters','CharactersWithSpaces']:e.text='0'
                b=etree.tostring(r,xml_declaration=True,encoding='UTF-8',standalone=True)
            elif n=='word/settings.xml':
                r=etree.fromstring(b);e=r.find(qn('w:updateFields'))
                if e is None:e=OxmlElement('w:updateFields');r.append(e)
                e.set(qn('w:val'),'true');b=etree.tostring(r,xml_declaration=True,encoding='UTF-8',standalone=True)
            out.writestr(n,b)
    tmp.unlink()

def pathology(d,c):
    para(d,'病理资料',11,True,heading=True)
    if not c['pathology_before']:
        para(d,'无候选时点前取材的独立病理报告记录；癌症登记中的组织学与在先NGS癌种编码见上表。');return
    para(d,'以下病理均为取材在先，报告可用时间未确认。检测“否”表示未检测，不是阴性结果。',9)
    for i,p in enumerate(c['pathology_before'],1):
        e=p['event'];f=e['facts'];r=p['raw']['fields']
        para(d,f"病理 {i}  D{e['event_day']}  {trans(f['procedure_type'])}  {trans(f['procedure'])}",10,True,heading=True)
        pr=para(d,f"报告编号 {r.get('path_proc_number')}/{r.get('path_rep_number')}；取材 {trans(r.get('path_proc_inst'))}；报告 {trans(r.get('path_rep_inst'))}；{loc(e['source'])}。",9)
        pr.paragraph_format.keep_with_next=True
        data=[]
        for x in f['specimens']:
            detail=f"浸润性肿瘤：{trans(x['invasive'])}；原位癌：{trans(x['in_situ'])}"
            if x['invasive_histology']:detail+='\n浸润性组织学：'+x['invasive_histology']
            if x['in_situ_histology']:detail+='\n原位组织学：'+x['in_situ_histology']
            if x['cancer_type']:detail+='\n癌种：'+trans(x['cancer_type'])
            data.append([str(x['specimen_number']),x['site'],detail])
        table(d,['标本','取材部位','病理记录'],data,[.5,2.25,4.3])
        tests=[]
        for prefix,title in [('pdl1','PD-L1'),('msi','MSI')]:
            testing=r.get(prefix+'_testing','')
            tests.append(title+'检测：'+trans(testing))
            if testing=='Yes':
                tests.extend(k+'='+v for k,v in r.items() if k.startswith(prefix) and v and k!=prefix+'_testing')
        para(d,'；'.join(tests)+'。',9)

def case_info(d,c):
    cf=c['cancer']['fields'];pf=c['patient']['fields'];ng=c['ngs_reports'][0]
    para(d,c['case_id']+' 病例资料',14,True,heading=True,page=True)
    table(d,None,[['诊断年龄',cf['age_dx']+' 岁','性别',trans(pf.get('naaccr_sex_code',''))],
        ['原发部位',cf['ca_d_site'],'初诊分期',trans(cf['stage_dx'])],
        ['登记组织学',trans(cf['ca_histology']),'候选时点',f"D{c['t0_day']}"],
        ['NGS报告日',f"D{ng['available_day']}",'癌种编码',ng['facts']['cpt_oncotree_code']]], [1.05,2.47,1.05,2.48],True)
    para(d,'患者编号 '+c['patient_key'][1]+'；癌症序号 '+c['cancer_seq']+'；'+SCOPE.get(c['scope_at_t0'],c['scope_at_t0'])+'。')
    para(d,'D 为相对首次索引癌症诊断的天数。登记来源：'+loc(c['cancer']['source'])+'；候选编号 '+c['candidate_id']+'。',9)
    pathology(d,c)
    para(d,'既往系统治疗',11,True,heading=True)
    rs=[e for e in c['regimens_before'] if e['cancer_seq']==c['cancer_seq']]
    if rs:table(d,['开始日','候选日前已记录的组分及日期','来源'],[[f"D{e['event_day']}",drugs(e['facts']['components'],True),loc(e['source'])] for e in rs],[.7,4.65,1.7])
    else:para(d,'未记录同癌症既往系统治疗。')
    other=[e for e in c['regimens_before'] if e['cancer_seq']!=c['cancer_seq']]
    if other:para(d,'其他癌症或归属不明的既往治疗另列于资料附录，未计入本癌症治疗史。')
    para(d,'候选日前最近评估',11,True,heading=True)
    recent=[]
    for kind in ['imaging','oncology_assessment','tumor_marker_collection']:
        recent.extend([e for e in c['events_before'] if e['kind']==kind][-2:])
    if recent:table(d,['日期与类型','原记录摘要','来源'],[[f"D{e['event_day']}\n{TYPES.get(e['kind'],e['kind'])}",facttext(e),loc(e['source'])] for e in sorted(recent,key=lambda e:e['event_day'])],[1.05,4.3,1.7])
    else:para(d,'无候选日前的相应评估记录。')
    para(d,'检查或就诊日期在先；未记录独立签发日期的资料仍保留可用时间未确认状态。完整事件及字段见资料索引。',9)
    rad=[e for e in c['events_before'] if e['kind']=='radiation_start' and e['cancer_seq']==c['cancer_seq']]
    if rad:
        para(d,'在先放疗开始记录：'+'；'.join(f"D{e['event_day']} {e['facts'].get('rt_type','')}，{loc(e['source'])}" for e in rad)+'。全疗程剂量和分次另存于资料索引，不据开始日期认定已经完成。')
    para(d,'在先分子检测',11,True,heading=True)
    for e in c['ngs_reports']:
        f=e['facts'];para(d,f"D{e['available_day']}；样本 {f['cpt_genie_sample_id']}；面板 {f['cpt_seq_assay_id']}；样本类型 {f['sample_type']}。{loc(e['source'])}。")
    for s in c['molecular_samples']:
        muts=s['mutations'];groups={}
        for m in muts:groups.setdefault(m['fields']['Hugo_Symbol'],[]).append(m['fields']['HGVSp_Short'] or m['fields']['HGVSc'])
        para(d,'小变异共 '+str(len(muts))+' 条。'+('；'.join(k+' '+', '.join(v) for k,v in groups.items()) if muts else '小变异表无该样本记录，不等同于阴性。'))
        cn=[g+'='+v['raw_value'] for g,v in s['cna_values'].items() if v['raw_value'] not in ['0','NA','']]
        para(d,'非零离散拷贝数记录：'+('；'.join(cn) if cn else '无非零记录')+'。数值保持数据集原编码；不能直接替代蛋白表达或临床靶向结论。')
        for sv in s['structural_variants']:
            f=sv['fields'];para(d,'结构变异：'+f['Event_Info']+'；DNA支持 '+f['DNA_Support']+'；RNA支持 '+f['RNA_Support']+'。'+f['Comments']+' '+loc(sv['source']))
    para(d,'详细变异位置、转录本、原始注释和来源见资料附录。基因或变异出现本身不代表致病性、胚系身份或可用药结论。',9)

def build_main():
    d=new('晚期胰腺癌候选治疗方案','Pilot 标签独立审核表')
    para(d,'请逐例核对病例与证据，并为每个候选方案选择一个标签；同一病例可有多个获得支持的方案。预拟标签用于审核，均可修改。',11)
    para(d,'审核专家编号 __________    姓名 __________    日期 __________',11)
    para(d,'填写说明',12,True,heading=True,page=True)
    para(d,'本表供两名专家分别填写。独立审核完成前不交换结果；分歧记录在裁定表。审核对象为病例时点与候选方案的证据匹配关系，不是疗效预测或方案优先级。')
    table(d,['标签','判定含义','填写要求'],[
        ['支持匹配','现有允许使用的病例事实满足所引路径的关键条件，未发现实质性冲突。','注明支持依据。'],
        ['条件性匹配','已有明确证据路径；有限、具体的关键条件尚待确认，条件成立后可匹配。','写明所需条件。'],
        ['有依据不匹配','已知病例事实与所评价路径的必要条件存在明确冲突。','指出冲突双方来源。'],
        ['无法判定','方案身份、资料或证据不足，或实质性矛盾尚无法解释。','说明缺口影响何项判断。']], [1.15,4.3,1.6])
    para(d,'研究性方案单独标明，不因有论文而自动获得支持。四类标签不是等级分数；未填写也不等同于无法判定。已使用的方案与其他候选按相同规则审核。')
    para(d,'资料读取',11,True,heading=True)
    para(d,'每例依次列出初诊资料、病理标本、在先疗程、最近评估及NGS。全部病理标本均保留，包括阴性标本；癌症登记的组织学不自动改写为明确导管亚型。病理取材日与报告可用日分别处理，NGS报告须严格早于候选日。')
    para(d,'时间未确认的信息不能改写为当时已知。候选日后加药、后续疗效与生存记录放在资料附录，不能用于提高或降低本时点标签。主表评估病理、分子、疾病情境和治疗史的证据关系，不替代完整的个体用药评估。')
    para(d,'证据编号对应附录中的来源、公开日期和定位。目录采用已核验的公开指南及原始研究；最新版CSCO与NCCN完整推荐表尚未逐项核验。发现候选遗漏或不同适用路径时，请写入病例末尾的补充栏。')
    para(d,'病例编号 PILOT-01 至 PILOT-25；方案编号为病例编号后加 T01 等。可使用 Word 查找编号定位病例和审核记录。',10)
    for c in cases:
        case_info(d,c)
        para(d,'候选方案标签审核',12,True,heading=True)
        cp=[p for p in pairs if p['case_id']==c['case_id']]
        for p in cp:
            para(d,p['pair_id']+'  '+p['regimen_display_name'],11,True,heading=True)
            facts='适用路径：'+p['assessment_context']+'\n判定理由：'+p['rationale']
            if p['unknown_conditions']:facts+='\n待核条件：'+'；'.join(p['unknown_conditions'])
            if p['unmet_conditions']:facts+='\n已知冲突：'+'；'.join(p['unmet_conditions'])
            facts+='\n证据：'+('；'.join(p['evidence_ids']) or '完整方案身份或相应组成证据尚未定位')
            t=table(d,['项目','内容','医生判定／备注'],[['预拟标签',p['proposed_label_zh']+'\n'+CTX[p['use_context']],''],['证据判断',facts,''],['医生判定','□ 支持匹配    □ 条件性匹配\n□ 有依据不匹配    □ 无法判定',''],['修改理由','适用条件、事实或证据的补充：\n','']], [1.0,4.5,1.55])
            for cell in t.rows[3].cells:shade(cell,'FFF8C6')
            # Keep a short review table together; long rows remain expandable.
            for row in t.rows[:-1]:
                for cell in row.cells:
                    for par in cell.paragraphs:par.paragraph_format.keep_with_next=True
        para(d,'病例补充意见',11,True,heading=True)
        para(d,'病理或分子解释修订／候选方案遗漏／其他使用路径：\n__________________________________________________________________\n__________________________________________________________________')
    save(d,'pilot_label_review_v2.0.docx')

def build_adjudication():
    d=new('晚期胰腺癌候选治疗方案','Pilot 标签分歧裁定表')
    para(d,'在两名专家完成独立审核后汇总。保留双方原始标签与理由；不得用预拟标签替代未提交的专家意见。最终标签、裁定理由与签名均由审核人员填写。')
    para(d,'专家一 __________    专家二 __________    裁定人 __________    日期 __________')
    para(d,'标签代码：S 支持匹配；C 条件性匹配；M 有依据不匹配；I 无法判定。空白表示尚未完成。分歧既包括主标签不同，也包括适用路径、研究属性或关键条件不同。')
    for c in cases:
        para(d,c['case_id']+' 裁定记录',14,True,heading=True,page=True)
        para(d,c['patient_key'][1]+f"    候选时点 D{c['t0_day']}")
        cp=[p for p in pairs if p['case_id']==c['case_id']]
        table(d,['方案编号与名称','专家一／专家二','最终标签与裁定理由'],[[p['pair_id']+'\n'+p['regimen_display_name'],'专家一：____\n专家二：____','最终标签：____\n理由／保留条件：\n'] for p in cp],[2.6,1.3,3.15])
        para(d,'研究属性或方案范围修订：________________________________________________')
        para(d,'尚未解决的问题：________________________________________________________')
        para(d,'裁定人签名：______________________    日期：______________________')
    save(d,'pilot_adjudication_v2.0.docx')

def build_reference():
    d=new('晚期胰腺癌候选治疗方案','Pilot 证据与病例资料附录')
    para(d,'第一部分为主审核表的证据来源；第二部分为分子明细、未定时的登记病理字段及实际治疗随访参考。后续事件不属于候选时点输入。完整事件、原始字段与来源在同目录的资料索引中查询。')
    para(d,'证据来源',14,True,heading=True,page=True)
    for e in sources:
        para(d,e['evidence_id']+'  '+e['title'],11,True,heading=True)
        para(d,'公开日期：'+e['publication_date']+'；来源类型：'+e['source_type']+'；读取层级：'+e['verification_level']+'。',9)
        para(d,e['finding_zh'])
        para(d,'定位：'+e['locator']+'\n'+e['url'],9)
    for c in cases:
        para(d,c['case_id']+' 资料附录',14,True,heading=True,page=True)
        para(d,c['patient_key'][1]+f"    候选时点 D{c['t0_day']}")
        cf=c['cancer']['fields']
        para(d,'登记中的病理分期与分级',11,True,heading=True)
        fields=['ca_grade','ca_path_t_stage','ca_path_n_stage','ca_path_group_stage','ca_clin_t_stage','ca_clin_n_stage','summary_stage','ca_tx_pre_path_stage','naaccr_tnm_edition_num']
        names=['分级','病理T','病理N','病理分期组','临床T','临床N','汇总分期','病理分期前治疗','TNM版本']
        para(d,'这些字段来自癌症级登记，无独立记录日期；不可自动视为候选日已知。')
        table(d,['字段','原登记值','来源'],[[n,val(cf[k]),loc(c['cancer']['source'])] for k,n in zip(fields,names)],[1.4,3.95,1.7])
        para(d,'小变异完整记录',11,True,heading=True)
        mm=[m for s in c['molecular_samples'] for m in s['mutations']]
        if mm:
            table(d,['基因与蛋白改变','核酸与坐标及注释','来源行'],[[m['fields']['Hugo_Symbol']+' '+val(m['fields']['HGVSp_Short']),m['fields']['HGVSc']+'\n'+m['fields']['NCBI_Build']+' chr'+m['fields']['Chromosome']+':'+m['fields']['Start_Position']+' '+m['fields']['Reference_Allele']+'>'+m['fields']['Tumor_Seq_Allele2']+'\n'+m['fields']['Variant_Classification']+'；原状态 '+m['fields']['Mutation_Status'],str(m['source']['physical_line'])] for m in mm],[1.75,4.5,.8])
            para(d,'来源 data_mutations_extended.txt；逐条原始字段、测序深度和原始注释见资料索引。原状态 SOMATIC 不等于已经过独立的致病性或胚系判定。',9)
        else:para(d,'该样本在小变异表中没有记录。')
        para(d,'实际治疗及后续参考',11,True,heading=True)
        para(d,'本节在标签独立填写后用于历史核查；结局不能证明候选时点匹配，也不能比较未观察到的替代方案疗效。')
        ca=c['candidate'];para(d,'候选日药物：'+drugs(ca['observed_action_drugs']))
        para(d,'完整疗程组分及开始日：'+'；'.join(drugs([x])+f" D{x.get('start_day')}" for x in ca['regimen_components_full_course']))
        op=next(p for p in pairs if p['case_id']==c['case_id'] and p['observed_regimen'])
        historical=[EI[x] for x in op['evidence_ids'] if 'phase_' in EI[x]['source_type']]
        para(d,'历史研究来源：'+('；'.join(x['evidence_id']+'（'+x['publication_date']+'）' for x in historical) or '本组成的历史原始研究尚未定位')+'。病例具体历年与当时可得性未确认，不能据公开日期推定同期适用。')
        follow=[e for e in c['events_after'] if e['kind'] in ['imaging','oncology_assessment']][:4]
        if follow:table(d,['日期与类型','候选日后的首批评估','来源'],[[f"D{e['event_day']}\n{TYPES[e['kind']]}",facttext(e),loc(e['source'])] for e in follow],[1.05,4.3,1.7])
        else:para(d,'未记录候选日后的影像或肿瘤科评估。')
        para(d,'后续完整疗程、全部评估、肿瘤标志物、生存字段及其他癌症记录在资料索引中单独保留。')
    save(d,'pilot_reference_appendix_v2.0.docx')

def build_html():
    esc=lambda x:html.escape(str(x))
    def render(x):
        if isinstance(x,dict):return '<dl>'+''.join('<dt>'+esc(k)+'</dt><dd>'+render(v)+'</dd>' for k,v in x.items())+'</dl>'
        if isinstance(x,list):return '<ol>'+''.join('<li>'+render(v)+'</li>' for v in x)+'</ol>'
        return esc(x) if x not in ['',None] else '<span class="muted">未记录</span>'
    out=['<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>Pilot 病例资料索引 v2.0</title><style>body{font:16px/1.6 Microsoft YaHei,sans-serif;max-width:1100px;margin:36px auto;color:#222;padding:0 24px}summary{cursor:pointer;color:#365f91;font-weight:600}details{margin:14px 0}dt{font-weight:600;overflow-wrap:anywhere}dd{margin:0 0 8px 24px;overflow-wrap:anywhere}ol{padding-left:24px}.muted{color:#777}nav{display:flex;flex-wrap:wrap;gap:14px}h1,h2{color:#111}section{border-top:1px solid #ddd;margin-top:30px}a{color:#365f91}th,td{padding:8px;border:1px solid #ddd;vertical-align:top}table{border-collapse:collapse;width:100%}</style><h1>Pilot 病例资料索引 v2.0</h1><p>本索引用于逐字段追溯，含候选时点前、同日、之后及时间不明的记录。后续资料与全程登记字段不能整体用作候选时点输入。每项原始记录保留来源文件、行号和内容。</p><nav>']
    for c in cases:out.append('<a href="#'+c['case_id']+'">'+c['case_id']+'</a>')
    out.append('</nav>')
    for c in cases:
        out.append('<section id="'+c['case_id']+'"><h2>'+c['case_id']+' '+esc(c['patient_key'][1])+'</h2><p>候选时点 D'+str(c['t0_day'])+'</p>')
        for k,title in [('pathology_before','在先取材病理报告全部字段'),('regimens_before','按候选日截断的既往疗程'),('molecular_samples','在先NGS关联样本全部字段'),('events_before','候选日之前事件'),('events_same_day','候选日同日事件'),('events_after','候选日之后事件'),('events_unknown','日期不明事件')]:
            out.append('<details><summary>'+title+'（'+str(len(c[k]))+'）</summary>')
            for i,x in enumerate(c[k],1):out.append('<details><summary>记录 '+str(i)+'</summary>'+render(x)+'</details>')
            out.append('</details>')
        out.append('<details><summary>全程临床原表及结局字段</summary>')
        for k,rs in c['all_clinical_records'].items():
            out.append('<details><summary>'+esc(k)+'（'+str(len(rs))+'）</summary>')
            for r in rs:out.append('<details><summary>'+esc(loc(r['source']))+'</summary>'+render(r)+'</details>')
            out.append('</details>')
        out.append('</details></section>')
    (O/'pilot_case_source_index_v2.0.html').write_text(''.join(out)+'</html>',encoding='utf8')

if __name__=='__main__':
    O.mkdir(parents=True,exist_ok=True)
    build_main();build_adjudication();build_reference();build_html()
    print('Built 3 DOCX files and complete source index.')
