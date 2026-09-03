#!/usr/bin/env python3
"""Generate the anonymous Chinese Pilot reviewer package with pre-t0 NGS evidence."""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import sys
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import audit_cohort_lock_label_feasibility as cohort_audit  # noqa: E402
import build_pilot_evidence_labels as pilot_builder  # noqa: E402
import privacy_checks as privacy  # noqa: E402


MUTATION_GENES = ("BRAF", "BRCA1", "BRCA2", "ERBB2")
FUSION_GENES = {"NRG1", "NTRK1", "NTRK2", "NTRK3", "RET"}
JSON_FIELDS = (
    "mutation_coverage_genes",
    "mutation_noncoverage_genes",
    "candidate_relevant_variants",
    "fusion_events",
    "msi_results",
    "mmr_results",
    "interpretation_notes",
)
EVIDENCE_FIELDS = [
    "pilot_case_code",
    "ngs_availability",
    "report_recency_band",
    "specimen_category",
    "dna_panel_gene_count",
    "mutation_coverage_genes",
    "mutation_noncoverage_genes",
    "candidate_relevant_variants",
    "fusion_events",
    "erbb2_cna_status",
    "msi_mmr_evidence_state",
    "msi_results",
    "mmr_results",
    "tmb_status",
    "her2_ihc_status",
    "germline_brca_status",
    "interpretation_notes",
]
RECENCY_LABELS = {
    "within_30_days": "距 t0 不超过30天",
    "31_to_90_days": "距 t0 31-90天",
    "91_to_180_days": "距 t0 91-180天",
    "181_to_365_days": "距 t0 181-365天",
    "more_than_365_days": "距 t0 超过365天",
    "unknown": "时间间隔未知",
}
MSI_MMR_LABELS = {
    "positive_msi_h_or_dmmr": "MSI-H或dMMR阳性",
    "explicit_nonpositive_requires_manual_review": "存在非阳性结果，需复核组合判定",
    "indeterminate_requires_manual_review": "结果不确定，需人工复核",
    "unknown_no_pre_t0_linked_result": "无可关联的t0前结果",
}


def load_yaml(path: Path) -> dict[str, Any]:
    value = pilot_builder.load_yaml(path)
    if not isinstance(value, dict):
        raise ValueError(f"YAML root is not a mapping: {path}")
    return value


def _clean_text(value: Any, limit: int = 240) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if not text or text.lower() in {"nan", "nat", "none"}:
        return ""
    if privacy.contains_identifier(text):
        return "[已隐去潜在标识信息]"
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def report_recency_band(t0_day: Any, report_day: Any) -> str:
    t0 = pilot_builder.numeric(t0_day)
    report = pilot_builder.numeric(report_day)
    if t0 is None or report is None or report >= t0:
        return "unknown"
    delta = t0 - report
    if delta <= 30:
        return "within_30_days"
    if delta <= 90:
        return "31_to_90_days"
    if delta <= 180:
        return "91_to_180_days"
    if delta <= 365:
        return "181_to_365_days"
    return "more_than_365_days"


def normalize_specimen(value: Any) -> str:
    text = _clean_text(value).lower()
    if text == "primary tumor":
        return "primary_tumor"
    if "metast" in text:
        return "metastatic_tumor_site_unspecified"
    return "other_or_unspecified"


def load_gene_panels(raw_root: Path) -> dict[str, set[str]]:
    panels: dict[str, set[str]] = {}
    for path in sorted(raw_root.rglob("data_gene_panel_*.txt")):
        stable_id = ""
        genes: set[str] = set()
        for raw_line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
            if raw_line.startswith("stable_id:"):
                stable_id = raw_line.split(":", 1)[1].strip()
            elif raw_line.startswith("gene_list:"):
                genes = {value.strip().upper() for value in raw_line.split("\t")[1:] if value.strip()}
        if stable_id:
            panels[stable_id] = genes
    return panels


def mutation_record(row: pd.Series) -> dict[str, Any]:
    depth = pilot_builder.numeric(row.get("t_depth"))
    alt = pilot_builder.numeric(row.get("t_alt_count"))
    vaf = ""
    if depth is not None and alt is not None and depth > 0:
        vaf = round(100.0 * alt / depth, 1)
    gene = _clean_text(row.get("Hugo_Symbol"), 40).upper()
    protein = _clean_text(row.get("HGVSp_Short"), 80)
    flag = "candidate_gene_variant_requires_review"
    if gene == "BRAF" and protein.lower().replace("p.", "") == "v600e":
        flag = "braf_v600e_candidate_signal"
    elif gene == "BRAF":
        flag = "braf_non_v600e"
    elif gene in {"BRCA1", "BRCA2"}:
        flag = "tumor_brca_not_germline"
    elif gene == "ERBB2":
        flag = "erbb2_sequence_variant_not_her2_ihc"
    return {
        "gene": gene,
        "coding_change": _clean_text(row.get("HGVSc"), 100),
        "protein_change": protein,
        "variant_classification": _clean_text(row.get("Variant_Classification"), 80),
        "consequence": _clean_text(row.get("Consequence"), 100),
        "mutation_status": _clean_text(row.get("Mutation_Status"), 40),
        "vaf_percent": vaf,
        "tumor_depth": int(depth) if depth is not None else "",
        "interpretation_flag": flag,
    }


def fusion_record(row: pd.Series) -> dict[str, Any]:
    gene_1 = _clean_text(row.get("Site1_Hugo_Symbol"), 40).upper()
    gene_2 = _clean_text(row.get("Site2_Hugo_Symbol"), 40).upper()
    matched = sorted({gene_1, gene_2} & FUSION_GENES)
    return {
        "gene_pair": "--".join(value for value in (gene_1, gene_2) if value),
        "candidate_fusion_genes": matched,
        "sv_class": _clean_text(row.get("Class"), 60),
        "breakpoint_type": _clean_text(row.get("Breakpoint_Type"), 60),
        "dna_support": _clean_text(row.get("DNA_Support"), 40),
        "rna_support": _clean_text(row.get("RNA_Support"), 40),
        "event_info": _clean_text(row.get("Event_Info"), 120),
        "annotation": _clean_text(row.get("Annotation"), 180),
        "interpretation_flag": "candidate_fusion_requires_annotation_review" if matched else "other_fusion_background",
    }


def discrete_cna_status(value: Any) -> str:
    parsed = pilot_builder.numeric(value)
    if parsed is None:
        return "unavailable"
    mapping = {
        -2.0: "deep_deletion",
        -1.0: "shallow_deletion",
        0.0: "no_discrete_cna_alteration_recorded",
        1.0: "low_level_gain",
        2.0: "high_level_amplification",
    }
    return mapping.get(parsed, "unrecognized_discrete_value")


def _json(value: Iterable[Any]) -> str:
    return json.dumps(list(value), ensure_ascii=False, separators=(",", ":"))


def _load_selected_pilot(repo_root: Path) -> tuple[pd.DataFrame, dict[str, Any]]:
    protocol = load_yaml(repo_root / "code" / "config" / "pilot_label_protocol_v0.1.yaml")
    context, _ = pilot_builder.build_decision_context(repo_root)
    selected = pilot_builder.assign_pilot_strata(context, protocol)
    linkage_path = repo_root / protocol["outputs"]["private_directory"] / protocol["outputs"]["private_case_linkage"]
    linkage = pd.read_csv(linkage_path, dtype=str, keep_default_na=False)
    expected = selected[["pilot_case_code", "index_ngs_sample_id"]].astype(str).sort_values("pilot_case_code")
    observed = linkage[["pilot_case_code", "index_ngs_sample_id"]].astype(str).sort_values("pilot_case_code")
    if expected.reset_index(drop=True).to_dict("records") != observed.reset_index(drop=True).to_dict("records"):
        raise ValueError("current deterministic Pilot does not match the stored private linkage")
    return selected, protocol


def build_reviewer_evidence(repo_root: Path) -> list[dict[str, Any]]:
    selected, _ = _load_selected_pilot(repo_root)
    raw_root = repo_root / cohort_audit.PANC_RELATIVE_ROOT
    clinical = pd.read_csv(
        pilot_builder.find_raw_file(raw_root, "data_clinical_sample.txt"),
        sep="\t",
        comment="#",
        dtype=str,
        keep_default_na=False,
    )
    matrix = pd.read_csv(
        pilot_builder.find_raw_file(raw_root, "data_gene_matrix.txt"),
        sep="\t",
        dtype=str,
        keep_default_na=False,
    )
    mutations = pd.read_csv(
        pilot_builder.find_raw_file(raw_root, "data_mutations_extended.txt"),
        sep="\t",
        dtype=str,
        low_memory=False,
    )
    structural = pd.read_csv(
        pilot_builder.find_raw_file(raw_root, "data_sv.txt"),
        sep="\t",
        dtype=str,
        low_memory=False,
    )
    cna = pd.read_csv(
        pilot_builder.find_raw_file(raw_root, "data_CNA.txt"),
        sep="\t",
        dtype=str,
        low_memory=False,
        index_col=0,
    )
    panels = load_gene_panels(raw_root)
    clinical_by_sample = clinical.set_index("SAMPLE_ID", drop=False)
    matrix_by_sample = matrix.set_index("SAMPLE_ID", drop=False)
    rows: list[dict[str, Any]] = []

    for _, context in selected.sort_values("pilot_case_code").iterrows():
        sample_id = str(context["index_ngs_sample_id"])
        if not sample_id or sample_id not in clinical_by_sample.index or sample_id not in matrix_by_sample.index:
            raise ValueError(f"selected pre-t0 NGS metadata is unavailable for {context['pilot_case_code']}")
        sample_meta = clinical_by_sample.loc[sample_id]
        panel_meta = matrix_by_sample.loc[sample_id]
        if isinstance(sample_meta, pd.DataFrame) or isinstance(panel_meta, pd.DataFrame):
            raise ValueError(f"duplicate NGS metadata for {context['pilot_case_code']}")
        mutation_panel = str(panel_meta.get("mutations", ""))
        panel_genes = panels.get(mutation_panel, set())
        covered = [gene for gene in MUTATION_GENES if gene in panel_genes]
        not_covered = [gene for gene in MUTATION_GENES if gene not in panel_genes]

        sample_mutations = mutations.loc[
            mutations["Tumor_Sample_Barcode"].astype(str).eq(sample_id)
            & mutations["Hugo_Symbol"].astype(str).isin(MUTATION_GENES)
        ].copy()
        variant_rows = [mutation_record(row) for _, row in sample_mutations.iterrows()]

        sample_sv = structural.loc[structural["Sample_Id"].astype(str).eq(sample_id)].copy()
        if not sample_sv.empty:
            narrative = sample_sv[["Event_Info", "Annotation", "Comments"]].fillna("").apply(
                lambda row: " ".join(row.astype(str)), axis=1
            ).str.lower()
            sample_sv = sample_sv.loc[narrative.str.contains("fusion", regex=False)]
        fusion_rows = [fusion_record(row) for _, row in sample_sv.iterrows()]

        erbb2_cna = "unavailable"
        if "ERBB2" in cna.index and sample_id in cna.columns:
            erbb2_cna = discrete_cna_status(cna.at["ERBB2", sample_id])

        msi_values, _ = pilot_builder._pre_t0_result_values(context, "msi", "msi_result")
        mmr_values, _ = pilot_builder._pre_t0_result_values(context, "mmr", "mmr_result")
        notes = [
            "Only the selected NGS report strictly before t0 is summarized.",
            "An absent variant call is not a clinical negative result without assay coverage and source confirmation.",
            "Fusion absence is not interpretable as fusion-negative from the available structured data.",
            "Tumor BRCA findings do not establish germline BRCA status.",
            "ERBB2 copy number does not establish HER2 IHC 3+ status.",
        ]
        rows.append(
            {
                "pilot_case_code": str(context["pilot_case_code"]),
                "ngs_availability": "available_strictly_pre_t0",
                "report_recency_band": report_recency_band(context["t0_start_day"], context["index_ngs_report_day"]),
                "specimen_category": normalize_specimen(sample_meta.get("SAMPLE_TYPE_DETAILED")),
                "dna_panel_gene_count": len(panel_genes),
                "mutation_coverage_genes": _json(covered),
                "mutation_noncoverage_genes": _json(not_covered),
                "candidate_relevant_variants": _json(variant_rows),
                "fusion_events": _json(fusion_rows),
                "erbb2_cna_status": erbb2_cna,
                "msi_mmr_evidence_state": str(context["msi_mmr_evidence_state"]),
                "msi_results": _json(_clean_text(value, 180) for value in msi_values),
                "mmr_results": _json(_clean_text(value, 180) for value in mmr_values),
                "tmb_status": "unavailable_unknown",
                "her2_ihc_status": "unavailable_unknown",
                "germline_brca_status": "unavailable_unknown",
                "interpretation_notes": _json(notes),
            }
        )
    return rows


def validate_reviewer_evidence(rows: list[dict[str, Any]]) -> None:
    if not rows:
        raise ValueError("reviewer NGS evidence is empty")
    if set(rows[0]) != set(EVIDENCE_FIELDS):
        raise ValueError("reviewer NGS evidence schema mismatch")
    frame = pd.DataFrame(rows, columns=EVIDENCE_FIELDS)
    if frame["pilot_case_code"].duplicated().any():
        raise ValueError("reviewer NGS evidence contains duplicate case codes")
    if any(privacy.is_high_risk_field_name(field) for field in frame.columns):
        raise ValueError("reviewer NGS evidence contains a high-risk identifier column")
    if privacy.contains_identifier(frame.to_csv(index=False)):
        raise ValueError("reviewer NGS evidence contains a GENIE identifier")
    if set(frame["ngs_availability"]) != {"available_strictly_pre_t0"}:
        raise ValueError("reviewer NGS evidence includes a non-pre-t0 assay")
    for field in JSON_FIELDS:
        for value in frame[field]:
            parsed = json.loads(value)
            if not isinstance(parsed, list):
                raise ValueError(f"{field} must encode a JSON list")


def write_evidence_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=EVIDENCE_FIELDS)
        writer.writeheader()
        writer.writerows(rows)


def _set_cell_width(cell: Any, width_inches: float) -> None:
    width = int(width_inches * 1440)
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def _set_table_geometry(table: Any, widths: list[float]) -> None:
    width_dxa = [int(width * 1440) for width in widths]
    table_grid = table._tbl.tblGrid
    for child in list(table_grid):
        table_grid.remove(child)
    for width in width_dxa:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(width))
        table_grid.append(grid_column)

    table_properties = table._tbl.tblPr
    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.insert(0, table_width)
    table_width.set(qn("w:w"), str(sum(width_dxa)))
    table_width.set(qn("w:type"), "dxa")
    table_indent = table_properties.find(qn("w:tblInd"))
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        table_properties.append(table_indent)
    table_indent.set(qn("w:w"), "0")
    table_indent.set(qn("w:type"), "dxa")


def _set_cell_margins(cell: Any, top: int = 80, start: int = 90, bottom: int = 80, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _shade(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _format_cell(cell: Any, *, bold: bool = False, color: str = "222222", size: float = 8.5) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(size)
            run.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)


def _set_repeat_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _prevent_row_split(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _insert_paragraph_after(document: DocumentObject, anchor: Any, text: str, *, heading: bool = False) -> Any:
    paragraph = document.add_paragraph()
    paragraph._p.getparent().remove(paragraph._p)
    anchor.addnext(paragraph._p)
    paragraph.paragraph_format.space_before = Pt(4 if heading else 2)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5 if heading else 9)
    run.bold = heading
    run.font.color.rgb = RGBColor(0x2F, 0x75, 0xB5) if heading else RGBColor(0x22, 0x22, 0x22)
    return paragraph


def _insert_table_after(
    document: DocumentObject,
    anchor: Any,
    rows: list[list[str]],
    widths: list[float],
    *,
    header: bool,
    label_columns: set[int] | None = None,
) -> Any:
    table = document.add_table(rows=len(rows), cols=len(widths))
    table.style = "Table Grid"
    table.autofit = False
    _set_table_geometry(table, widths)
    table._tbl.getparent().remove(table._tbl)
    anchor.addnext(table._tbl)
    label_columns = label_columns or set()
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = value
            _set_cell_width(cell, widths[column_index])
            is_header = header and row_index == 0
            is_label = column_index in label_columns
            if is_header:
                _shade(cell, "2F75B5")
                _format_cell(cell, bold=True, color="FFFFFF", size=8.5)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif is_label:
                _shade(cell, "D9EAF7")
                _format_cell(cell, bold=True, color="1F4E79", size=8.5)
            else:
                _format_cell(cell)
        _prevent_row_split(table.rows[row_index])
    if header:
        _set_repeat_header(table.rows[0])
    return table


def _decode_list(row: dict[str, Any], field: str) -> list[Any]:
    value = row.get(field, "[]")
    parsed = json.loads(value) if isinstance(value, str) else value
    return parsed if isinstance(parsed, list) else []


def _specimen_label(value: str) -> str:
    return {
        "primary_tumor": "原发灶",
        "metastatic_tumor_site_unspecified": "转移灶（部位未细分）",
        "other_or_unspecified": "其他或未说明",
    }.get(value, "未说明")


def _cna_label(value: str) -> str:
    return {
        "high_level_amplification": "高水平扩增信号",
        "low_level_gain": "低水平增加",
        "no_discrete_cna_alteration_recorded": "未记录离散拷贝数改变",
        "shallow_deletion": "浅缺失",
        "deep_deletion": "深缺失",
        "unavailable": "不可用",
        "unrecognized_discrete_value": "结果编码需复核",
    }.get(value, "不可用")


def molecular_table_rows(evidence: dict[str, Any]) -> list[list[str]]:
    rows = [["类别", "结果", "证据明细", "解释边界"]]
    variants = _decode_list(evidence, "candidate_relevant_variants")
    if variants:
        for variant in variants:
            result = " ".join(value for value in (variant.get("gene", ""), variant.get("protein_change", "")) if value)
            detail_parts = [
                variant.get("coding_change", ""),
                variant.get("variant_classification", ""),
                variant.get("mutation_status", ""),
            ]
            if variant.get("vaf_percent") != "":
                detail_parts.append(f"VAF {variant['vaf_percent']}%")
            if variant.get("tumor_depth") != "":
                detail_parts.append(f"深度 {variant['tumor_depth']}×")
            boundary = {
                "braf_v600e_candidate_signal": "用于BRAF V600E候选条件复核",
                "braf_non_v600e": "非V600E，不能满足该候选标志物",
                "tumor_brca_not_germline": "肿瘤体细胞信号，不能证明胚系BRCA",
                "erbb2_sequence_variant_not_her2_ihc": "序列变异不能替代HER2 IHC 3+",
            }.get(variant.get("interpretation_flag"), "候选相关性需专家复核")
            rows.append(["小变异", result or "变异记录", "；".join(value for value in detail_parts if value), boundary])
    else:
        rows.append(["小变异", "未见候选相关变异记录", "结合上表DNA检测覆盖理解", "无记录不等于临床阴性"])

    fusions = _decode_list(evidence, "fusion_events")
    if fusions:
        for fusion in fusions:
            support = f"DNA支持：{fusion.get('dna_support') or '未知'}；RNA支持：{fusion.get('rna_support') or '未知'}"
            annotation = "；".join(
                value
                for value in (
                    fusion.get("sv_class", ""),
                    fusion.get("breakpoint_type", ""),
                    fusion.get("event_info", ""),
                    fusion.get("annotation", ""),
                    support,
                )
                if value
            )
            boundary = "需确认报告注释为功能性融合并符合候选条件"
            if fusion.get("interpretation_flag") == "other_fusion_background":
                boundary = "非冻结候选池特异标志物，仅作背景"
            rows.append(["融合/SV", fusion.get("gene_pair") or "融合记录", annotation, boundary])
    else:
        rows.append(["融合/SV", "未见明确融合记录", "结构化数据未提供可验证的阴性覆盖结论", "不能据此判定融合阴性"])

    rows.append(
        [
            "拷贝数",
            f"ERBB2：{_cna_label(str(evidence['erbb2_cna_status']))}",
            "离散CNA结构化结果",
            "不能替代HER2 IHC 3+或完整HER2检测",
        ]
    )
    msi = _decode_list(evidence, "msi_results")
    mmr = _decode_list(evidence, "mmr_results")
    pathology_detail = "；".join([*(f"MSI：{value}" for value in msi), *(f"MMR：{value}" for value in mmr)])
    rows.append(
        [
            "MSI/MMR",
            MSI_MMR_LABELS.get(str(evidence["msi_mmr_evidence_state"]), "结果需复核"),
            pathology_detail or "无可关联的t0前原始结果",
            "按MSI与MMR组合条件判定",
        ]
    )
    rows.extend(
        [
            ["胚系BRCA", "当前数据不可用", "肿瘤NGS不能确定胚系状态", "奥拉帕利维持条件仍需独立证据"],
            ["TMB", "当前数据不可用", "未提取可靠TMB数值和单位", "保持未知"],
            ["HER2 IHC", "当前数据不可用", "未获得IHC 3+证据", "保持未知"],
        ]
    )
    return rows


def append_ngs_sections(document: DocumentObject, evidence_rows: list[dict[str, Any]]) -> None:
    headings = [paragraph for paragraph in document.paragraphs if re.match(r"^PILOT-\d{3}", paragraph.text.strip())]
    summary_tables = [
        table
        for table in list(document.tables)
        if len(table.rows) == 4 and table.cell(0, 0).text.strip() == "既往晚期方案数"
    ]
    evidence_rows = sorted(evidence_rows, key=lambda row: row["pilot_case_code"])
    heading_codes = [re.match(r"^(PILOT-\d{3})", paragraph.text.strip()).group(1) for paragraph in headings]
    evidence_codes = [str(row["pilot_case_code"]) for row in evidence_rows]
    if len(summary_tables) != len(evidence_rows) or heading_codes != evidence_codes:
        raise ValueError("template case order does not match reviewer evidence")

    for summary_table, evidence in zip(summary_tables, evidence_rows):
        anchor = summary_table._tbl
        heading = _insert_paragraph_after(document, anchor, "t0前分子检测证据", heading=True)
        covered = "、".join(_decode_list(evidence, "mutation_coverage_genes")) or "无"
        uncovered = "、".join(_decode_list(evidence, "mutation_noncoverage_genes")) or "无"
        overview = [
            ["NGS资料", "已关联严格早于t0的结构化检测", "报告时距", RECENCY_LABELS[str(evidence["report_recency_band"])]],
            ["标本类型", _specimen_label(str(evidence["specimen_category"])), "DNA panel", f"约{evidence['dna_panel_gene_count']}个基因"],
            ["小变异覆盖", covered, "未覆盖", uncovered],
            ["融合阴性范围", "来源数据不足以确认", "输出范围", "仅列冻结候选池相关证据"],
        ]
        overview_table = _insert_table_after(
            document,
            heading._p,
            overview,
            [1.05, 2.45, 1.05, 2.55],
            header=False,
            label_columns={0, 2},
        )
        _insert_table_after(
            document,
            overview_table._tbl,
            molecular_table_rows(evidence),
            [0.75, 1.45, 3.15, 1.75],
            header=True,
        )


def _replace_in_runs(paragraph: Any, old: str, new: str) -> bool:
    changed = False
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            changed = True
    return changed


def _replace_paragraph_text(paragraph: Any, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def update_package_front_matter(document: DocumentObject) -> None:
    note_anchor = None
    for paragraph in document.paragraphs:
        _replace_in_runs(paragraph, "v1.0", "v1.1")
        if paragraph.text.strip().startswith("晚期 PDAC 候选治疗标签，我们出了"):
            _replace_paragraph_text(
                paragraph,
                "本次Pilot共24个匿名决策点×16个主池候选方案=384个病例-候选对。两名专家独立审核，共768次审核；按每次10元计，每名专家3840元，合计7680元。",
            )
        if "信息不足时选择" in paragraph.text:
            note_anchor = paragraph._p
    if note_anchor is None:
        raise ValueError("could not locate reviewer instruction block in template")
    note = _insert_paragraph_after(
        document,
        note_anchor,
        "分子证据说明：各病例仅展示t0前选定NGS样本的候选池相关结构化字段。无变异记录不等于检测阴性；肿瘤BRCA不等于胚系BRCA；ERBB2拷贝数不替代HER2 IHC 3+；无融合记录不等于融合阴性。",
    )
    note.paragraph_format.space_before = Pt(4)
    document.core_properties.title = "晚期PDAC候选治疗标签Pilot专家审核包 v1.1"
    document.core_properties.subject = "Track A证据与约束标签双专家复核"


def document_text(document: DocumentObject) -> str:
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def validate_document(document: DocumentObject, expected_cases: int) -> None:
    text = document_text(document)
    if privacy.contains_identifier(text):
        raise ValueError("reviewer package contains a GENIE identifier")
    case_codes = set(re.findall(r"PILOT-\d{3}", text))
    if len(case_codes) != expected_cases:
        raise ValueError("reviewer package case coverage mismatch")
    if "v1.1" not in text or text.count("t0前分子检测证据") != expected_cases:
        raise ValueError("reviewer package NGS sections are incomplete")
    if "本次Pilot共24个匿名决策点×16个主池候选方案=384个病例-候选对" not in text:
        raise ValueError("reviewer package front-matter counts are inconsistent")


def build_document(template: Path, output: Path, rows: list[dict[str, Any]]) -> None:
    if template.resolve() == output.resolve():
        raise ValueError("output must not overwrite the reviewer-package template")
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, output)
    document = Document(output)
    update_package_front_matter(document)
    append_ngs_sections(document, rows)
    validate_document(document, len(rows))
    document.save(output)
    validate_document(Document(output), len(rows))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo-root", type=Path, default=Path("."))
    parser.add_argument("--template", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--evidence-output", type=Path)
    args = parser.parse_args()
    repo_root = args.repo_root.resolve()
    protocol = load_yaml(repo_root / "code" / "config" / "pilot_label_protocol_v0.1.yaml")
    output_config = protocol["outputs"]
    private_dir = repo_root / output_config["private_directory"]
    template = args.template or private_dir / output_config["private_reviewer_package_template"]
    output = args.output or private_dir / output_config["private_reviewer_package"]
    evidence_output = args.evidence_output or private_dir / output_config["private_reviewer_ngs_evidence"]

    rows = build_reviewer_evidence(repo_root)
    validate_reviewer_evidence(rows)
    write_evidence_csv(evidence_output, rows)
    build_document(template, output, rows)
    print(
        json.dumps(
            {
                "decision_points": len(rows),
                "reviewer_ngs_evidence": str(evidence_output),
                "reviewer_package": str(output),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
