from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from docx import Document


ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "code" / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

import generate_pilot_reviewer_package as package


def synthetic_evidence() -> dict[str, object]:
    row = {field: "" for field in package.EVIDENCE_FIELDS}
    row.update(
        {
            "pilot_case_code": "PILOT-001",
            "ngs_availability": "available_strictly_pre_t0",
            "report_recency_band": "31_to_90_days",
            "specimen_category": "primary_tumor",
            "dna_panel_gene_count": 300,
            "mutation_coverage_genes": '["BRAF","BRCA1","BRCA2","ERBB2"]',
            "mutation_noncoverage_genes": "[]",
            "candidate_relevant_variants": "[]",
            "fusion_events": "[]",
            "erbb2_cna_status": "no_discrete_cna_alteration_recorded",
            "msi_mmr_evidence_state": "unknown_no_pre_t0_linked_result",
            "msi_results": "[]",
            "mmr_results": "[]",
            "tmb_status": "unavailable_unknown",
            "her2_ihc_status": "unavailable_unknown",
            "germline_brca_status": "unavailable_unknown",
            "interpretation_notes": "[]",
        }
    )
    return row


def test_report_recency_band_is_strictly_pre_t0():
    assert package.report_recency_band(100, 99) == "within_30_days"
    assert package.report_recency_band(100, 69) == "31_to_90_days"
    assert package.report_recency_band(100, 100) == "unknown"
    assert package.report_recency_band(100, 101) == "unknown"


def test_mutation_record_keeps_reviewable_fields_without_identifiers():
    result = package.mutation_record(
        pd.Series(
            {
                "Hugo_Symbol": "BRCA2",
                "HGVSc": "c.100del",
                "HGVSp_Short": "p.A34fs",
                "Variant_Classification": "Frame_Shift_Del",
                "Consequence": "frameshift_variant",
                "Mutation_Status": "SOMATIC",
                "t_alt_count": "12",
                "t_depth": "48",
            }
        )
    )
    assert result["vaf_percent"] == 25.0
    assert result["interpretation_flag"] == "tumor_brca_not_germline"
    assert not {"sample_id", "patient_id", "record_id"} & set(result)


def test_discrete_cna_status_is_not_her2_interpretation():
    assert package.discrete_cna_status("2") == "high_level_amplification"
    assert package.discrete_cna_status("0") == "no_discrete_cna_alteration_recorded"
    assert package.discrete_cna_status("") == "unavailable"


def test_reviewer_evidence_schema_is_anonymous():
    package.validate_reviewer_evidence([synthetic_evidence()])
    assert not any(package.privacy.is_high_risk_field_name(field) for field in package.EVIDENCE_FIELDS)


def test_ngs_sections_are_inserted_after_case_summary():
    document = Document()
    document.add_table(rows=5, cols=3).cell(0, 0).text = "标签"
    document.add_paragraph("PILOT-001 匿名决策点 1/1")
    summary = document.add_table(rows=4, cols=4)
    summary.cell(0, 0).text = "既往晚期方案数"
    document.add_paragraph("本决策点复核重点")
    candidate = document.add_table(rows=17, cols=4)
    candidate.cell(0, 0).text = "序号"

    package.append_ngs_sections(document, [synthetic_evidence()])

    assert len(document.tables) == 5
    assert package.document_text(document).count("t0前分子检测证据") == 1
    assert "无记录不等于临床阴性" in package.document_text(document)
    assert [column.get(package.qn("w:w")) for column in document.tables[2]._tbl.tblGrid] == ["1512", "3528", "1512", "3671"]
